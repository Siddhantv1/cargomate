import io

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter

from app.ingest import MAX_BYTES, MAX_CHARS, IngestError, ingest, to_chunks

DOC = b"EPCG scheme note.\n\nZero customs duty on capital goods with an export obligation."


def put(store, embed, name="a.txt", data=DOC, cat="dgft", day="2026-04-01"):
    return ingest(store, embed, name, data, cat, day)


def test_create_unchanged_update(store, embed):
    assert put(store, embed)[0] == "created"
    assert put(store, embed)[0] == "unchanged"
    assert put(store, embed, data=DOC + b"\n\nNew paragraph.")[0] == "updated"
    assert [d["name"] for d in store.list_docs()] == ["a.txt"]


def test_reupload_replaces_old_chunks(store, embed):
    put(store, embed, data=b"Old rate is five percent.")
    put(store, embed, data=b"New rate is seven percent.")
    texts = [h["text"] for h in store.search(embed(["rate percent"])[0], min_score=0)]
    assert texts == ["New rate is seven percent."]


def test_metadata_change_alone_triggers_update(store, embed):
    put(store, embed, day="2026-04-01")
    assert put(store, embed, day="2026-07-01")[0] == "updated"
    assert store.list_docs()[0]["effective_date"] == "2026-07-01"


def test_failed_embedding_keeps_old_version(store, embed):
    put(store, embed, data=b"Original content.")

    def boom(_):
        raise RuntimeError("openai down")

    with pytest.raises(RuntimeError):
        put(store, boom, data=b"Replacement content.")
    assert [h["text"] for h in store.search(embed(["content"])[0], min_score=0)] == ["Original content."]


def test_csv_rows_keep_their_column_names():
    csv = b"mode,route,rate\nair,DEL-DXB,95\nsea,NSA-JEA,12\n"
    chunks = to_chunks("r.csv", csv)
    assert len(chunks) == 1
    assert "mode: air; route: DEL-DXB; rate: 95" in chunks[0]
    assert "mode: sea; route: NSA-JEA; rate: 12" in chunks[0]


def test_csv_with_bom_and_blank_cells():
    assert to_chunks("r.csv", "\ufeffa,b\n1,\n,\n".encode("utf-8")) == ["a: 1"]


def test_xlsx_rows_get_sheet_prefix():
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["lane", "days"])
    ws.append(["DEL-DXB", 2])
    buf = io.BytesIO()
    wb.save(buf)
    assert to_chunks("r.xlsx", buf.getvalue()) == ["[Rates] lane: DEL-DXB; days: 2"]


def test_docx_paragraphs_and_tables():
    from docx import Document

    d = Document()
    d.add_paragraph("Checklist intro.")
    t = d.add_table(rows=2, cols=2)
    for i, v in enumerate(["doc", "needed"]):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(["invoice", "yes"]):
        t.rows[1].cells[i].text = v
    buf = io.BytesIO()
    d.save(buf)
    chunks = to_chunks("c.docx", buf.getvalue())
    assert "Checklist intro." in chunks[0]
    assert "doc: invoice; needed: yes" in chunks[-1]


def test_long_paragraph_is_split_to_size():
    chunks = to_chunks("big.txt", ("word " * 3000).encode())
    assert len(chunks) > 1 and all(len(c) <= MAX_CHARS for c in chunks)


def test_paragraphs_are_packed_not_one_per_chunk():
    assert to_chunks("p.txt", b"one\n\ntwo\n\nthree") == ["one\n\ntwo\n\nthree"]


@pytest.mark.parametrize(
    "kw, msg",
    [
        (dict(name="a.exe"), "Unsupported file type"),
        (dict(cat="nope"), "Unknown category"),
        (dict(day="01/04/2026"), "YYYY-MM-DD"),
        (dict(data=b"   \n\n  "), "No text found"),
        (dict(data=b"x" * (MAX_BYTES + 1)), "larger than"),
    ],
)
def test_validation_errors(store, embed, kw, msg):
    with pytest.raises(IngestError, match=msg):
        put(store, embed, **kw)
    assert store.list_docs() == []


def test_blank_pdf_reports_scanned_pdf(store, embed):
    w = PdfWriter()
    w.add_blank_page(100, 100)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(IngestError, match="OCR"):
        put(store, embed, name="scan.pdf", data=buf.getvalue())


def test_corrupt_file_is_an_ingest_error(store, embed):
    with pytest.raises(IngestError, match="Could not read"):
        put(store, embed, name="bad.pdf", data=b"not a pdf")


def test_path_in_filename_is_stripped(store, embed):
    put(store, embed, name="C:\\Users\\me\\..\\notes.txt")
    assert store.list_docs()[0]["name"] == "notes.txt"
