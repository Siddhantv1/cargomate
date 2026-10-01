"""Parse -> chunk -> embed -> atomically swap a document's chunks."""
import csv
import hashlib
import io
import os
import re
from datetime import date

from .store import CATEGORIES, Store

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 3000  # ~750 tokens
EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx")


class IngestError(ValueError):
    """The upload is unusable; the message is safe to show the admin."""


def _pack(parts: list[str], sep: str) -> list[str]:
    chunks, cur = [], ""
    for p in parts:
        while len(p) > MAX_CHARS:  # one oversized part: hard split
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(p[:MAX_CHARS])
            p = p[MAX_CHARS:]
        if not p:
            continue
        if cur and len(cur) + len(sep) + len(p) > MAX_CHARS:
            chunks.append(cur)
            cur = ""
        cur = f"{cur}{sep}{p}" if cur else p
    if cur:
        chunks.append(cur)
    return chunks


def _prose(text: str) -> list[str]:
    return _pack([p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()], "\n\n")


def _table(rows: list[list], prefix: str = "") -> list[str]:
    """Each row becomes 'header: value; ...' so a rate never loses its column names when chunked."""
    rows = [[str(c).strip() if c is not None else "" for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if len(rows) < 2:
        return []
    header, lines = rows[0], []
    for r in rows[1:]:
        line = "; ".join(f"{h}: {v}" for h, v in zip(header, r) if v)
        if line:
            lines.append(prefix + line)
    return _pack(lines, "\n")


def _pdf(data: bytes) -> list[str]:
    from pypdf import PdfReader

    return _prose("\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages))


def _docx(data: bytes) -> list[str]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    chunks = _prose("\n\n".join(parts))
    for t in doc.tables:
        chunks += _table([[c.text for c in row.cells] for row in t.rows])
    return chunks


def _xlsx(data: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks = []
    for ws in wb.worksheets:
        chunks += _table([list(r) for r in ws.iter_rows(values_only=True)], prefix=f"[{ws.title}] ")
    return chunks


def to_chunks(name: str, data: bytes) -> list[str]:
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf":
            chunks = _pdf(data)
        elif ext == ".docx":
            chunks = _docx(data)
        elif ext == ".xlsx":
            chunks = _xlsx(data)
        elif ext == ".csv":
            chunks = _table(list(csv.reader(io.StringIO(data.decode("utf-8-sig")))))
        else:
            chunks = _prose(data.decode("utf-8-sig"))
    except Exception as e:
        raise IngestError(f"Could not read {name}: {e}") from e
    if not chunks:
        raise IngestError(f"No text found in {name}. Scanned PDFs need OCR before upload.")
    return chunks


def ingest(store: Store, embed, filename: str, data: bytes, category: str, effective_date: str):
    """Returns (status, n_chunks); status is created | updated | unchanged."""
    name = os.path.basename(filename.replace("\\", "/")).strip()
    if not name.lower().endswith(EXTENSIONS):
        raise IngestError(f"Unsupported file type. Use one of: {', '.join(EXTENSIONS)}.")
    if category not in CATEGORIES:
        raise IngestError(f"Unknown category '{category}'. Use one of: {', '.join(CATEGORIES)}.")
    try:
        date.fromisoformat(effective_date)
    except ValueError:
        raise IngestError("Effective date must be YYYY-MM-DD.") from None
    if len(data) > MAX_BYTES:
        raise IngestError(f"File is larger than {MAX_BYTES // 1024 // 1024} MB.")

    fp = hashlib.sha256(data + f"|{category}|{effective_date}".encode()).hexdigest()
    old = store.fingerprint(name)
    if old == fp:
        return "unchanged", 0

    chunks = to_chunks(name, data)
    vectors = embed(chunks)  # raises before anything is written, so a failure keeps the old version
    store.upsert(name, fp, category, effective_date, chunks, vectors)
    return ("updated" if old else "created"), len(chunks)
